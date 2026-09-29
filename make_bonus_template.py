"""Qo'shimcha daromad % shabloni: bazadagi barcha brend × kategoriya juftliklari, savdo bo'yicha saralangan.
Amaldagi foizlar bo'lsa, ular ham to'ldiriladi.

    python make_bonus_template.py Qoshimcha_daromad_shablon.xlsx
"""
import asyncio
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from db import create_pool, migrate

SQL = """
SELECT c.name category, p.brand,
       (SELECT bb.pct FROM brand_bonus bb WHERE bb.brand = p.brand
          AND (bb.category_id = c.id OR bb.category_id IS NULL)
          ORDER BY bb.category_id NULLS LAST LIMIT 1) pct,
       count(DISTINCT p.product_id) skus,
       coalesce(sum(s.amount), 0) amount,
       min(s.sale_date) d1, max(s.sale_date) d2
FROM products p JOIN categories c ON c.id = p.category_id
LEFT JOIN sales_daily s ON s.product_id = p.product_id
GROUP BY c.id, c.name, p.brand
ORDER BY c.name, amount DESC, p.brand
"""


async def main(path: str) -> None:
    pool = await create_pool()
    await migrate(pool)
    rows = await pool.fetch(SQL)
    period = await pool.fetchrow("SELECT min(sale_date) d1, max(sale_date) d2 FROM sales_daily")
    await pool.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Бренд фоизлари"
    sales_hdr = f"Савдо {period['d1']:%d.%m}–{period['d2']:%d.%m}, млн" if period["d1"] else "Савдо, млн"
    ws.append(["Категория", "Бренд", "Қўшимча даромад %", "SKU сони", sales_hdr])
    head_fill = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    for r in rows:
        ws.append([r["category"], r["brand"], float(r["pct"]) if r["pct"] is not None else None,
                   r["skus"], round(float(r["amount"]) / 1e6, 1)])
        ws.cell(ws.max_row, 3).fill = input_fill
        ws.cell(ws.max_row, 3).number_format = "0.00"
        ws.cell(ws.max_row, 5).number_format = "#,##0.0"
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="100",
                        error="0 dan 100 gacha son kiriting (3.5 = 3.5%)", showErrorMessage=True)
    ws.add_data_validation(dv)
    dv.add(f"C2:C{max(ws.max_row, 2)}")
    for col, w in zip("ABCDE", (22, 18, 14, 10, 16)):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{ws.max_row}"

    info = wb.create_sheet("Изоҳ")
    for line in [
        "Қўшимча даромад % = доп доход + ретро + компенсация (жами битта фоиз).",
        "Фақат сариқ устунни тўлдиринг: 3.5 = 3.5%. Бўш қолса — 0%.",
        "Категория бўш бўлса, фоиз бренднинг барча категорияларига қўлланади.",
        "Қўшимча даромад (сўм) = Жами(Кирим нархи) × % ; Маржа = Валовка + қўшимча даромад.",
        "Файл ҳар юкланганда тўлиқ алмаштирилади ва бутун тарих янги фоиз билан ҳисобланади.",
        "SKU сони ва Савдо устунлари — маълумот учун, юкланмайди.",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 100
    wb.save(path)
    print(f"{path}: {len(rows)} ta brend × kategoriya")


if __name__ == "__main__":
    asyncio.run(main(sys.argv[1] if len(sys.argv) > 1 else "Qoshimcha_daromad_shablon.xlsx"))
