"""Bazadagi umumiy tovar spravochnigini bitta Excel'ga chiqarish (/spravochnik, Mini App → Yuklash).

Fayl tuzilishi botga yuklanadigan tovar spravochnigi bilan bir xil — tahrirlab, shu holicha qaytarsa bo'ladi:
  «Tovarlar» — Группа, Категория, Товар Ид, Товар номи, Бренд, Статус + barcha kategoriyalarning xususiyat
               ustunlari («Nomi (birlik)»). Har qatorda faqat o'z kategoriyasining ustunlari: bo'sh bo'lsa — sariq,
               boshqa kategoriyaniki — kulrang.
  «Kategoriyalar» — kategoriya bo'yicha SKU soni, xususiyatlar va to'ldirilmaganlar soni.
Faqat Mini App'da ko'rinadigan (is_filter) xususiyatlar chiqadi.
"""
from __future__ import annotations

import html
import io
import json
from datetime import date

import asyncpg
from openpyxl import Workbook

from export_missing import BOLD, HEAD_FILL, NA_FILL, NA_FONT, TODO_FILL, _autowidth


def _label(a) -> str:
    return f"{a['name']} ({a['unit']})" if a["unit"] else a["name"]


async def products_export(con: asyncpg.Connection) -> tuple[bytes, str, str] | None:
    """(xlsx, fayl nomi, xabar matni) yoki None — bazada tovar yo'q bo'lsa."""
    cats = await con.fetch("""SELECT c.id, c.grp, c.name, count(p.product_id) n
                              FROM categories c JOIN products p ON p.category_id = c.id
                              GROUP BY c.id ORDER BY c.grp, c.name""")
    if not cats:
        return None
    attrs: dict[int, list] = {}
    for a in await con.fetch("""SELECT category_id, slot, name, unit FROM category_attributes
                                WHERE is_filter ORDER BY category_id, sort_order, slot"""):
        attrs.setdefault(a["category_id"], []).append(a)
    cols: list[str] = []                                   # ustunlar — kategoriyalar tartibida, takrorsiz
    for c in cats:
        for a in attrs.get(c["id"], []):
            if _label(a) not in cols:
                cols.append(_label(a))
    prods = await con.fetch("""SELECT p.product_id, p.name, p.brand, p.status, p.attrs, p.category_id
                               FROM products p JOIN categories c ON c.id = p.category_id
                               ORDER BY c.grp, c.name, p.brand, p.name""")
    cat_by_id = {c["id"]: c for c in cats}

    wb = Workbook()
    ws = wb.active
    ws.title = "Tovarlar"
    head = ["Группа", "Категория", "Товар Ид", "Товар номи", "Бренд", "Статус"] + cols
    ws.append(head)
    for i in range(1, len(head) + 1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = BOLD, HEAD_FILL
    empty_by_cat: dict[int, int] = {}
    for r, p in enumerate(prods, start=2):
        c = cat_by_id[p["category_id"]]
        av = p["attrs"] if isinstance(p["attrs"], dict) else json.loads(p["attrs"] or "{}")
        own = {_label(a): av.get(f"a{a['slot']}") for a in attrs.get(c["id"], [])}
        ws.append([c["grp"], c["name"], p["product_id"], p["name"], p["brand"], p["status"]]
                  + [own.get(col) for col in cols])
        miss = False
        for i, col in enumerate(cols, start=7):
            cell = ws.cell(row=r, column=i)
            if col not in own:
                cell.fill, cell.font = NA_FILL, NA_FONT
            elif own[col] in (None, ""):
                cell.fill = TODO_FILL
                miss = True
        if miss:
            empty_by_cat[c["id"]] = empty_by_cat.get(c["id"], 0) + 1
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    _autowidth(ws)

    wi = wb.create_sheet("Kategoriyalar")
    wi.append(["Группа", "Категория", "SKU", "Xususiyatlar", "Xususiyati to'liq emas (SKU)"])
    for i in range(1, 6):
        cell = wi.cell(row=1, column=i)
        cell.font, cell.fill = BOLD, HEAD_FILL
    for c in cats:
        wi.append([c["grp"], c["name"], c["n"], ", ".join(_label(a) for a in attrs.get(c["id"], [])),
                   empty_by_cat.get(c["id"], 0)])
    _autowidth(wi)

    buf = io.BytesIO()
    wb.save(buf)
    today = date.today()
    fname = f"tovar_spravochnigi_{today:%d.%m.%Y}.xlsx"

    esc = lambda v: html.escape(str(v or ""), quote=False)  # noqa: E731
    n = lambda v: f"{v:,}".replace(",", " ")  # noqa: E731
    lines = [f"📦 <b>Tovar spravochnigi</b> · {today:%d.%m.%Y} holati",
             f"{len(cats)} ta kategoriya · {n(len(prods))} ta SKU"]
    for c in cats:
        lines.append(f"• {esc(c['name'])} — {n(c['n'])} SKU · {len(attrs.get(c['id'], []))} xususiyat")
    tot_empty = sum(empty_by_cat.values())
    if tot_empty:
        lines.append(f"⚠️ Xususiyati to'ldirilmagan: {n(tot_empty)} SKU (Excel'da sariq)")
    lines += ["", "Tahrirlab, shu faylni botga qaytarsangiz — spravochnik sifatida yuklanadi."]
    return buf.getvalue(), fname, "\n".join(lines)
