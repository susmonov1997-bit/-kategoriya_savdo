"""Savdo yuklanganda spravochnikda topilmagan SKU va filiallarni Excel shablonga chiqarish.

Fayl tuzilishi:
  1-varaq «Tovarlar»   — tovar spravochnigi formatida: Группа, Категория, Товар Ид, Товар номи, Бренд, Статус
                         + shu kategoriyalarning xususiyat ustunlari (bo'sh, sariq). To'ldirib, aynan shu faylni
                         botga yuborsa bo'ladi — tovar spravochnigi sifatida yuklanadi.
  2-varaq «Filiallar»  — (bo'lsa) Территория, Филиал, Кластер — filial spravochnigiga qo'shish uchun.
  3-varaq «Ma'lumot»   — har bir SKU/filial bo'yicha qator soni, dona, summa va qisqa yo'riqnoma.
"""
from __future__ import annotations

import io
import json

import asyncpg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from loaders_common import name_key

HEAD_FILL = PatternFill("solid", fgColor="DDE7F5")
TODO_FILL = PatternFill("solid", fgColor="FFF4CC")
NA_FILL = PatternFill("solid", fgColor="E7E7E7")       # boshqa kategoriyaning ustuni — to'ldirilmaydi
NA_FONT = Font(color="9A9A9A")
BOLD = Font(bold=True)


def _d(v):
    from datetime import date as _date
    return _date.fromisoformat(v) if isinstance(v, str) else v


def has_missing(report: dict | None) -> bool:
    return bool(report and (report.get("unknown_skus") or report.get("unknown_branches")))


async def report_of(con, upload_id: int) -> dict | None:
    raw = await con.fetchval("SELECT report FROM uploads WHERE id = $1 AND kind = 'sales'", upload_id)
    if raw is None:
        return None
    return json.loads(raw) if isinstance(raw, str) else raw


def file_name(report: dict) -> str:
    if report.get("cumulative"):
        return f"kutayotganlar_{report.get('asof', '')}.xlsx"
    d = report.get("dates") or []
    period = (d[0] if len(d) == 1 else f"{d[0]}_{d[-1]}") if d else "savdo"
    return f"spravochnikda_yoq_{period}.xlsx"


def _autowidth(ws, min_w: int = 8, max_w: int = 60) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(min_w, min(max_w, width + 2))


async def build_missing_xlsx(con: asyncpg.Connection, report: dict) -> bytes | None:
    skus = report.get("unknown_skus") or []
    brs = report.get("unknown_branches") or []
    if not skus and not brs:
        return None

    # kategoriya → guruh va xususiyat ustunlari (bazadagi nomlar bilan, birlik qavsda)
    cats = {name_key(r["name"]): r for r in await con.fetch("SELECT id, name, grp FROM categories")}
    owners = {r["name_key"]: r["grp"] for r in await con.fetch("SELECT name_key, grp FROM category_owner_list")}
    attr_rows = await con.fetch(
        "SELECT category_id, name, unit FROM category_attributes WHERE is_filter ORDER BY category_id, sort_order, slot")
    attrs_by_cat: dict[int, list[str]] = {}
    for a in attr_rows:
        attrs_by_cat.setdefault(a["category_id"], []).append(f"{a['name']} ({a['unit']})" if a["unit"] else a["name"])

    attr_cols: list[str] = []
    for s in skus:
        c = cats.get(name_key(s.get("category")))
        for col in attrs_by_cat.get(c["id"], []) if c else []:
            if col not in attr_cols:
                attr_cols.append(col)

    wb = Workbook()
    # --- 1. Tovarlar (spravochnik formatida) ---
    ws = wb.active
    ws.title = "Tovarlar"
    head = ["Группа", "Категория", "Товар Ид", "Товар номи", "Бренд", "Статус"] + attr_cols
    ws.append(head)
    own_cols: list[set[str]] = []                         # har qatorda o'z kategoriyasining ustunlari
    for s in sorted(skus, key=lambda x: (x.get("category") or "", -(x.get("amount") or 0))):
        k = name_key(s.get("category"))
        grp = (cats[k]["grp"] if k in cats else None) or owners.get(k)
        ws.append([grp, s.get("category"), s.get("product_id"), s.get("name"), s.get("brand"), None]
                  + [None] * len(attr_cols))
        own_cols.append(set(attrs_by_cat.get(cats[k]["id"], [])) if k in cats else set(attr_cols))
    for i, h in enumerate(head, 1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = BOLD, HEAD_FILL
    # sariq — shu qatorda to'ldiriladi (Статус + o'z kategoriyasining xususiyatlari); kulrang — kerak emas
    for r, own in enumerate(own_cols, start=2):
        ws.cell(row=r, column=6).fill = TODO_FILL
        for i, col in enumerate(attr_cols, start=7):
            cell = ws.cell(row=r, column=i)
            if col in own:
                cell.fill = TODO_FILL
            else:
                cell.fill, cell.font = NA_FILL, NA_FONT
                cell.value = None
    ws.freeze_panes = "C2"
    _autowidth(ws)

    # --- 2. Filiallar ---
    if brs:
        wb2 = wb.create_sheet("Filiallar")
        wb2.append(["Территория", "Филиал", "Кластер"])
        for b in brs:
            wb2.append([None, b.get("name"), None])
        for i in range(1, 4):
            c = wb2.cell(row=1, column=i)
            c.font, c.fill = BOLD, HEAD_FILL
        for r in range(2, wb2.max_row + 1):
            wb2.cell(row=r, column=1).fill = TODO_FILL
            wb2.cell(row=r, column=3).fill = TODO_FILL
        _autowidth(wb2)

    # --- 3. Ma'lumot ---
    wi = wb.create_sheet("Ma'lumot")
    d = report.get("dates") or []
    cum = bool(report.get("cumulative"))
    if cum:
        wi.append(["Barcha kutayotganlar", f"{report.get('asof', '')} holati (hamma yuklangan fayllardan)"])
    else:
        wi.append(["Savdo fayli", ", ".join(filter(None, [d[0] if d else None, d[-1] if len(d) > 1 else None]))])
    wi.append(["Spravochnikda yo'q SKU", len(skus), "summa", round(sum(s.get("amount") or 0 for s in skus))])
    wi.append(["Spravochnikda yo'q filial", len(brs), "summa", round(sum(b.get("amount") or 0 for b in brs))])
    wi.append([])
    wi.append(["Qanday ishlatiladi:"])
    wi.append(["1. «Tovarlar» varag'ida faqat SARIQ kataklarni to'ldiring (Статус va o'z kategoriyasining xususiyatlari); KULRANG — boshqa kategoriyaniki, bo'sh qoladi."])
    wi.append(["2. Faylni botga yuboring — tovar spravochnigi sifatida yuklanadi (faqat shu SKU'lar qo'shiladi)."])
    wi.append(["3. «Filiallar» varag'ini filial spravochnigingizga ko'chiring (Территория, Кластер) va yuklang."])
    wi.append(["4. Savdo faylini qayta yuklash shart emas — bu SKU/filiallar savdosi saqlab qo'yilgan va spravochnik yuklanishi bilan avtomatik qo'shiladi."])
    wi.append([])
    wi.append(["Tur", "Товар Ид / Филиал", "Nomi", "Категория", "Qator", "Dona", "Summa"]
              + (["Birinchi sana", "Oxirgi sana", "Necha kun sotilgan"] if cum else []))
    for c in wi[wi.max_row]:
        c.font, c.fill = BOLD, HEAD_FILL
    for s in sorted(skus, key=lambda x: -(x.get("amount") or 0)):
        wi.append(["SKU", s.get("product_id"), s.get("name"), s.get("category"), s.get("rows"),
                   s.get("qty"), round(s.get("amount") or 0)]
                  + ([_d(s.get("first")), _d(s.get("last")), s.get("days")] if cum else []))
    for b in sorted(brs, key=lambda x: -(x.get("amount") or 0)):
        wi.append(["Filial", b.get("name"), None, None, b.get("rows"), None, round(b.get("amount") or 0)]
                  + ([_d(b.get("first")), _d(b.get("last")), b.get("days")] if cum else []))
    for r in wi.iter_rows(min_row=1, max_row=1):
        r[0].font = BOLD
    for row in wi.iter_rows(min_row=12):
        if row[6].value is not None:
            row[6].number_format = "#,##0"
        if cum:
            for c in row[7:9]:
                c.number_format = "DD.MM.YYYY"
    wi.column_dimensions["A"].width = 26
    for col in "BCDEFGHIJ":
        wi.column_dimensions[col].width = 18
    wi.column_dimensions["C"].width = 45
    for row in wi.iter_rows(min_row=5, max_row=9):
        row[0].alignment = Alignment(wrap_text=False)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def send_missing(bot_token: str, chat_id: int, data: bytes, report: dict) -> None:
    """Excel'ni botdan foydalanuvchiga yuboradi (Mini App'dan chaqirilganda)."""
    from aiogram import Bot
    from aiogram.types import BufferedInputFile
    bot = Bot(bot_token)
    try:
        await bot.send_document(chat_id, BufferedInputFile(data, filename=file_name(report)),
                                caption=missing_caption(report), reply_markup=all_pending_kb())
    finally:
        await bot.session.close()


def missing_caption(report: dict) -> str:
    n_s, n_b = len(report.get("unknown_skus") or []), len(report.get("unknown_branches") or [])
    parts = []
    if n_s:
        parts.append(f"{n_s} ta SKU")
    if n_b:
        parts.append(f"{n_b} ta filial")
    return (f"📎 Spravochnikda yo'q: {', '.join(parts)}.\n"
            "«Tovarlar» varag'ini to'ldirib botga yuboring — ularning savdosi avtomatik qo'shiladi, "
            "savdo faylini qayta yuklash shart emas.")


# ---------------------------------------------------------------------------
# Barcha kutayotganlar (/kutayotgan): hamma yuklangan fayllardan hali spravochnikka qo'shilmagan SKU va filiallar
def all_pending_kb():
    from aiogram.types import InlineKeyboardButton, InlineKeyboardMarkup
    return InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="📎 Barcha kutayotganlar",
                                                                       callback_data="pend:all")]])


def _dm(d) -> str:
    return d.strftime("%d.%m") if d else ""


async def pending_report(con: asyncpg.Connection) -> dict:
    """sales_pending dan: spravochnikda yo'q SKU'lar va filiallar (to'planib borgan), summa bo'yicha."""
    from datetime import date as _date
    skus = [{"product_id": int(r["pid"]), "name": r["name"], "category": r["cat"], "brand": r["brand"],
             "rows": int(r["n"]), "qty": float(r["qty"]), "amount": float(r["amount"]),
             "first": r["d1"].isoformat(), "last": r["d2"].isoformat(), "days": int(r["days"])}
            for r in await con.fetch(
                """SELECT sp.product_id pid, (array_agg(sp.product_name ORDER BY sp.sale_date DESC))[1] name, max(sp.category) cat, max(sp.brand) brand,
                          count(*) n, sum(sp.qty) qty, sum(sp.amount) amount, min(sp.sale_date) d1, max(sp.sale_date) d2,
                          count(DISTINCT sp.sale_date) days
                   FROM sales_pending sp
                   WHERE NOT EXISTS (SELECT 1 FROM products p WHERE p.product_id = sp.product_id)
                   GROUP BY sp.product_id ORDER BY sum(sp.amount) DESC""")]
    brs = [{"name": r["name"], "rows": int(r["n"]), "amount": float(r["amount"]),
            "first": r["d1"].isoformat(), "last": r["d2"].isoformat(), "days": int(r["days"])}
           for r in await con.fetch(
               """SELECT max(sp.branch_name) name, count(*) n, sum(sp.amount) amount, min(sp.sale_date) d1,
                         max(sp.sale_date) d2, count(DISTINCT sp.sale_date) days
                  FROM sales_pending sp
                  WHERE NOT EXISTS (SELECT 1 FROM branches b WHERE b.name_key = sp.branch_key)
                  GROUP BY sp.branch_key ORDER BY sum(sp.amount) DESC""")]
    return {"cumulative": True, "asof": _date.today().strftime("%d.%m.%Y"),
            "unknown_skus": skus, "unknown_branches": brs}


def pending_text(rep: dict, top: int = 5) -> str:
    import html
    esc = lambda v: html.escape(str(v or ""), quote=False)  # noqa: E731
    mln = lambda v: (f"{v / 1e9:,.2f} mlrd" if abs(v) >= 1e9 else f"{v / 1e6:,.1f} mln").replace(",", " ").replace(".", ",")  # noqa: E731
    skus, brs = rep.get("unknown_skus") or [], rep.get("unknown_branches") or []
    if not skus and not brs:
        return "✅ Kutayotgan savdo yo'q — hamma SKU va filiallar spravochnikda bor."
    out = ["⏳ <b>Spravochnikda yo'qligi uchun kutayotgan savdo</b>",
           f"<i>hamma yuklangan fayllardan, {esc(rep.get('asof'))} holati</i>", ""]
    if skus:
        d1, d2 = min(s["first"] for s in skus), max(s["last"] for s in skus)
        out.append(f"SKU: <b>{len(skus)} ta</b> · {mln(sum(s['amount'] for s in skus))} so'm · "
                   f"{d1[8:10]}.{d1[5:7]} — {d2[8:10]}.{d2[5:7]}")
    if brs:
        out.append(f"Filial: <b>{len(brs)} ta</b> · {mln(sum(b['amount'] for b in brs))} so'm")
    if skus:
        out += ["", f"Eng katta {min(top, len(skus))} ta:"]
        for i, s in enumerate(skus[:top], 1):
            out.append(f"{i}. {esc(s['name'] or s['product_id'])} · {esc(s['category'])} — {mln(s['amount'])} ({s['days']} kun)")
    out += ["", "📎 Excel'ni to'ldirib botga yuboring — savdo avtomatik qo'shiladi."]
    return "\n".join(out)
