"""Qoldiq API (011_stock.sql) — Mini App «Qoldiq» tabi va boshqa tablardagi qoldiq qatorlari.

Asosiy qoida: qoldiq — holat. Filtr davrining oxirgi sanasidan oldingi (yoki o'sha) eng so'nggi qoldiq olinadi;
«necha kunga yetadi» = qoldiq ÷ (shu sanagacha 30 kunlik savdo / 30).
  * Filial: filial qoldig'i ÷ filial savdosi.
  * Ombor: ombor qoldig'i ÷ u ta'minlaydigan filiallar savdosi (filiallardagi qoldiq alohida ko'rsatiladi).
Yetishmovchilik darajasi (joy × kategoriya × brend):
  0 🔴 tugadi     — < 3 kunga (0 ham)
  1 🟠 tugayapti  — < 7 kunga
  2 🟡 kamaymoqda — < 15 kunga va oxirgi 7 kunlik savdo 30 kunlik o'rtachadan ≥ 30% yuqori (savdo tezlashgan)
«Necha kunga yetadi» (signal, xarita, tafsilot) — 30 kunlik va 7 kunlik sur'atdan qaysi biri kam bo'lsa.
Oldingi qoldiq (~bir hafta oldin) faqat «yangi» belgisi va «o'tgan hafta» ko'rsatkichi uchun (bo'lmasa ham ishlaydi).
Xavf (so'm/kun) = shu joy × kategoriya × brendning kunlik o'rtacha savdosi.

Endpointlar (Filters + qo'shimcha maydonlar):
  POST /api/stock/signals   signallar ro'yxati + hisoblagichlar
  POST /api/stock/map       xarita: qator (ombor | hudud | filial) × brend → kunga yetadi
  POST /api/stock/detail    joy × kategoriya × brend tafsiloti: tarix, SKU'lar, boshqa omborlardagi qoldiq
  POST /api/stock/overview  umumiy qoldiq: jami, omborlar, hududlar
  POST /api/stock/by        o'lchov bo'yicha qoldiq (kategoriya/brend/SKU/hudud/klaster/filial) — boshqa tablardagi qator
"""
from __future__ import annotations

import json
import math
from datetime import date, timedelta
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException, Request
from pydantic import Field

from api_auth import User, current_user
from api_queries import _attr_cols
from api_schemas import Filters

router = APIRouter(prefix="/api/stock")


def _M():
    """api_main — kech import (aylanma importdan qochish: bot api_stock ni api_main'dan oldin ham import qiladi)."""
    import api_main
    return api_main

WINDOW = 30            # kunlik o'rtacha savdo oynasi
LV_DAYS = (3, 7, 15)   # daraja chegaralari (kun)
ACCEL = 1.3            # «kamaymoqda»: oxirgi 7 kunlik savdo 30 kunlikdan ≥ 30% yuqori
MIN_Q7 = 3             # 7 kunda kamida shuncha dona sotilgan bo'lsa — sur'at hisobga olinadi (tasodifiy sotuv emas)


class StockReq(Filters):
    loc: Literal["all", "w", "b"] = "all"
    level: int | None = None                     # 0 | 1 | 2 | None (hammasi)
    limit: int = Field(60, ge=1, le=500)
    offset: int = Field(0, ge=0)
    rows: Literal["w", "region", "b"] = "w"      # xarita qatorlari
    wid: int | None = None                       # xarita: shu omborning filiallari
    map_category: int | None = None              # xarita kategoriyasi (ko'p kategoriyali doirada)
    loc_type: Literal["w", "b", "region", "all"] = "w"
    loc_id: int | None = None
    brand: str | None = None
    dim: str = "category"


# ---------------------------------------------------------------------------
# Sanalar
async def snapshots(con, date_to: date) -> dict:
    """Qoldiq sanalari: joriy (date_to gacha eng so'nggi) va solishtirma (~bir hafta oldingi)."""
    out = {}
    for k, t in (("b", "stock_branch"), ("w", "stock_wh")):
        d = await con.fetchval(f"SELECT max(snap_date) FROM {t} WHERE snap_date <= $1", date_to) \
            or await con.fetchval(f"SELECT max(snap_date) FROM {t}")
        p = None
        if d:
            p = await con.fetchval(f"SELECT max(snap_date) FROM {t} WHERE snap_date <= $1", d - timedelta(days=6)) \
                or await con.fetchval(f"SELECT max(snap_date) FROM {t} WHERE snap_date < $1", d)
        out[k], out["p" + k] = d, p
    last_sale = await con.fetchval("SELECT max(sale_date) FROM sales_fact")
    out["sales_to"] = last_sale
    return out


def _end(d: date | None, last_sale: date | None) -> date | None:
    if d is None:
        return None
    return min(d, last_sale) if last_sale else d


# ---------------------------------------------------------------------------
class _SQL:
    """Parametrli SQL yig'uvchi: mahsulot va joy filtrlari."""

    def __init__(self, f: Filters, cats: list[int], attrs: dict):
        self.f, self.cats, self.attrs, self.params = f, cats, attrs, []

    def p(self, v) -> str:
        self.params.append(v)
        return f"${len(self.params)}"

    def ctes(self) -> str:
        f = self.f
        pc, bc = [], []
        if f.brands:
            pc.append(f"brand = ANY({self.p(f.brands)}::text[])")
        if f.statuses:
            pc.append(f"COALESCE(status,'—') = ANY({self.p(f.statuses)}::text[])")
        if f.category_ids:
            pc.append(f"category_id = ANY({self.p(f.category_ids)}::int[])")
        if f.owners:
            pc.append(f"owner = ANY({self.p(f.owners)}::text[])")
        if f.product_ids:
            pc.append(f"product_id = ANY({self.p(f.product_ids)}::bigint[])")
        for slot, vals in f.attrs.items():
            if vals and slot in self.attrs:
                pc.append(f"d_a{int(slot)} = ANY({self.p(vals)}::text[])")
        if f.region_ids:
            bc.append(f"region_id = ANY({self.p(f.region_ids)}::int[])")
        if f.cluster_ids:
            bc.append(f"cluster_id = ANY({self.p(f.cluster_ids)}::int[])")
        if f.branch_ids:
            bc.append(f"id = ANY({self.p(f.branch_ids)}::int[])")
        cats = self.p(self.cats)
        return f"""
pd0 AS (
    SELECT p.product_id, p.name AS pname, p.brand, p.status, p.category_id, c.name AS cname,
           COALESCE(c.owner, '—') AS owner{_attr_cols(self.attrs)}
    FROM products p JOIN categories c ON c.id = p.category_id WHERE p.category_id = ANY({cats}::int[])
), pd AS (SELECT * FROM pd0 {'WHERE ' + ' AND '.join(pc) if pc else ''}
), bb AS (
    SELECT id AS branch_id, region_id, cluster_id, warehouse_id FROM branches
    {'WHERE ' + ' AND '.join(bc) if bc else ''}
)"""


async def _ctx(con, f: Filters, user: User):
    cats, attrs, _ = await _M().scope_ctx(con, f, user)
    return cats, attrs


# ---------------------------------------------------------------------------
# Asosiy jadvallar: joy × kategoriya × brend
async def _branch_rows(con, f, cats, attrs, d: date, end: date) -> list:
    """(branch_id, category_id, brand) → stock, cost, sales qty/amount (30 kun), OOS (filial×SKU)."""
    s = _SQL(f, cats, attrs)
    w = s.ctes()
    D, S, E = s.p(d), s.p(end - timedelta(days=WINDOW - 1)), s.p(end)
    return await con.fetch(f"""
WITH {w},
cov AS (SELECT DISTINCT branch_id FROM stock_branch WHERE snap_date = {D}),
st AS (SELECT s.branch_id, s.product_id, s.qty, s.cost_uzs FROM stock_branch s
       JOIN bb USING (branch_id) JOIN pd USING (product_id) WHERE s.snap_date = {D}),
sl AS (SELECT e.branch_id, e.product_id, sum(e.qty) q, sum(e.amount) a,
              sum(e.qty) FILTER (WHERE e.sale_date > {E}::date - 7) q7 FROM sales_fact e
       JOIN bb USING (branch_id) JOIN pd USING (product_id) JOIN cov USING (branch_id)
       WHERE e.sale_date BETWEEN {S} AND {E} GROUP BY 1, 2),
j AS (SELECT branch_id, product_id, COALESCE(st.qty, 0) sq, COALESCE(st.cost_uzs, 0) sv,
             COALESCE(sl.q, 0) q, COALESCE(sl.a, 0) a, COALESCE(sl.q7, 0) q7
      FROM st FULL JOIN sl USING (branch_id, product_id))
SELECT j.branch_id, pd.category_id, pd.brand, sum(sq)::float8 sq, sum(sv)::float8 sv, sum(q)::float8 q, sum(a)::float8 a,
       count(*) FILTER (WHERE sq <= 0 AND q > 0) oos, sum(q7)::float8 q7
FROM j JOIN pd USING (product_id) GROUP BY 1, 2, 3""", *s.params)


async def _wh_rows(con, f, cats, attrs, d: date, end: date) -> list:
    """(warehouse_id, category_id, brand) → ombor qoldig'i, $ qiymati, filiallar savdosi, OOS SKU (omborda 0)."""
    s = _SQL(f, cats, attrs)
    w = s.ctes()
    D, S, E = s.p(d), s.p(end - timedelta(days=WINDOW - 1)), s.p(end)
    return await con.fetch(f"""
WITH {w},
wf AS (SELECT DISTINCT warehouse_id FROM bb WHERE warehouse_id IS NOT NULL
       AND warehouse_id IN (SELECT DISTINCT warehouse_id FROM stock_wh WHERE snap_date = {D})),
st AS (SELECT s.warehouse_id, s.product_id, s.qty, s.value_usd FROM stock_wh s
       JOIN wf USING (warehouse_id) JOIN pd USING (product_id) WHERE s.snap_date = {D}),
sl AS (SELECT bb.warehouse_id, e.product_id, sum(e.qty) q, sum(e.amount) a,
              sum(e.qty) FILTER (WHERE e.sale_date > {E}::date - 7) q7 FROM sales_fact e
       JOIN bb USING (branch_id) JOIN pd USING (product_id) JOIN wf ON wf.warehouse_id = bb.warehouse_id
       WHERE e.sale_date BETWEEN {S} AND {E} GROUP BY 1, 2),
j AS (SELECT warehouse_id, product_id, COALESCE(st.qty, 0) sq, COALESCE(st.value_usd, 0) sv,
             COALESCE(sl.q, 0) q, COALESCE(sl.a, 0) a, COALESCE(sl.q7, 0) q7
      FROM st FULL JOIN sl USING (warehouse_id, product_id))
SELECT j.warehouse_id, pd.category_id, pd.brand, sum(sq)::float8 sq, sum(sv)::float8 sv, sum(q)::float8 q, sum(a)::float8 a,
       count(*) FILTER (WHERE sq <= 0 AND q > 0) oos, sum(q7)::float8 q7
FROM j JOIN pd USING (product_id) GROUP BY 1, 2, 3""", *s.params)


def _days(stock: float, q: float) -> float | None:
    """Necha kunga yetadi; savdo bo'lmasa None (∞)."""
    if q <= 0:
        return None
    return stock / (q / WINDOW)


def _days_eff(stock: float, q: float, q7: float) -> float | None:
    """Necha kunga yetadi — 30 kunlik va 7 kunlik sur'atdan qaysi biri KAM bo'lsa (savdo tezlashsa — tezroq tugaydi)."""
    d30 = _days(stock, q)
    if q7 >= MIN_Q7:
        d7 = stock / (q7 / 7)
        return d7 if d30 is None else min(d30, d7)
    return d30


def _accel(q: float, q7: float) -> float | None:
    """Oxirgi 7 kunlik kunlik savdo 30 kunlik o'rtachadan necha marta yuqori (1,3 = +30%). Kam savdoda — None."""
    if q <= 0 or q7 < MIN_Q7:
        return None
    return (q7 / 7) / (q / WINDOW)


def _level(days: float | None, accel: float | None) -> int | None:
    if days is None:
        return None
    if days < LV_DAYS[0]:
        return 0
    if days < LV_DAYS[1]:
        return 1
    if days < LV_DAYS[2] and accel is not None and accel >= ACCEL:
        return 2
    return None


async def base_data(con, f: Filters, cats: list[int], attrs: dict) -> dict:
    """Joriy va oldingi qoldiq bo'yicha asosiy jadvallar (API va bot uchun umumiy)."""
    sn = await snapshots(con, f.date_to)
    ls = sn["sales_to"]
    out = {"snap": {k: (v.isoformat() if v else None) for k, v in sn.items()}, "b": [], "w": [], "pb": {}, "pw": {},
           "rate": 0.0}
    out["meta"] = {
        "branches": {r["id"]: dict(r) for r in await con.fetch(
            """SELECT b.id, b.name, b.region_id, b.cluster_id, b.warehouse_id, rg.name region
               FROM branches b JOIN regions rg ON rg.id = b.region_id""")},
        "wh": {r["id"]: r["name"] for r in await con.fetch("SELECT id, name FROM warehouses")},
        "regions": {r["id"]: r["name"] for r in await con.fetch("SELECT id, name FROM regions")},
        "cats": {r["id"]: r["name"] for r in await con.fetch("SELECT id, name FROM categories WHERE id = ANY($1::int[])", cats)},
    }
    out["links"] = {"total": len(out["meta"]["branches"]),
                    "unlinked": sum(1 for b in out["meta"]["branches"].values() if not b["warehouse_id"])}
    if sn["b"]:
        out["b"] = [tuple(r) for r in await _branch_rows(con, f, cats, attrs, sn["b"], _end(sn["b"], ls))]
        if sn["pb"]:
            out["pb"] = {(r[0], r[1], r[2]): _days_eff(r[3], r[5], r[8])
                         for r in await _branch_rows(con, f, cats, attrs, sn["pb"], _end(sn["pb"], ls))}
    if sn["w"]:
        out["w"] = [tuple(r) for r in await _wh_rows(con, f, cats, attrs, sn["w"], _end(sn["w"], ls))]
        if sn["pw"]:
            out["pw"] = {(r[0], r[1], r[2]): _days_eff(r[3], r[5], r[8])
                         for r in await _wh_rows(con, f, cats, attrs, sn["pw"], _end(sn["pw"], ls))}
    if sn["w"] or sn["b"]:
        out["rate"] = float(await con.fetchval(
            "SELECT rate FROM fx_periods WHERE $1 >= valid_from AND $1 < valid_to", sn["w"] or sn["b"]) or 0)
    return out


async def _base(request: Request, user: User, f: Filters) -> dict:
    """base_data — foydalanuvchi doirasida, keshlanadi."""
    async def run():
        async with request.app.state.pool.acquire() as con:
            cats, attrs = await _ctx(con, f, user)
            return await base_data(con, f, cats, attrs)
    return await _M().cached(request, user, "stock_base", f.model_dump(mode="json", include=set(Filters.model_fields)), run)


def _signals(base: dict, loc: str) -> list[dict]:
    meta, out = base["meta"], []
    br = meta["branches"]
    # filiallar OOS va qoldig'ini ombor × kategoriya × brendga yig'ish
    wb: dict[tuple, list] = {}
    for b, c, brand, sq, sv, q, a, oos, q7 in base["b"]:
        wid = br.get(b, {}).get("warehouse_id")
        if wid:
            x = wb.setdefault((wid, c, brand), [0.0, 0.0, 0])
            x[0] += sq
            x[1] += q
            x[2] += oos
    if loc in ("all", "b"):
        for b, c, brand, sq, sv, q, a, oos, q7 in base["b"]:
            d = _days_eff(sq, q, q7)
            prev = base["pb"].get((b, c, brand)) if base["pb"] else None
            ac = _accel(q, q7)
            lv = _level(d, ac)
            if lv is None:
                continue
            bi = br.get(b, {})
            out.append({"t": "b", "id": b, "name": bi.get("name"), "sub": bi.get("region"), "cat_id": c,
                        "cat": meta["cats"].get(c), "brand": brand, "level": lv, "days": d, "prev": prev,
                        "stock": sq, "sales_day": q / WINDOW, "risk": a / WINDOW, "oos": oos,
                        "accel": ac})
    if loc in ("all", "w"):
        for wid, c, brand, sq, sv, q, a, oos, q7 in base["w"]:
            d = _days_eff(sq, q, q7)
            prev = base["pw"].get((wid, c, brand)) if base["pw"] else None
            ac = _accel(q, q7)
            lv = _level(d, ac)
            if lv is None:
                continue
            bsq, bq, boos = wb.get((wid, c, brand), [0.0, 0.0, 0])
            out.append({"t": "w", "id": wid, "name": meta["wh"].get(wid), "sub": None, "cat_id": c,
                        "cat": meta["cats"].get(c), "brand": brand, "level": lv, "days": d, "prev": prev,
                        "stock": sq, "sales_day": q / WINDOW, "risk": a / WINDOW, "oos": boos, "oos_sku": oos, "accel": ac,
                        "br_stock": bsq, "br_days": _days(bsq, q)})
    for x in out:   # «yangi» — oldingi qoldiqda signal yo'q edi (15 kundan ko'pga yetardi)
        has_prev = bool(base["pb"] if x["t"] == "b" else base["pw"])
        x["new"] = has_prev and (x["prev"] is None or x["prev"] >= LV_DAYS[2])
    return out


def _r(v, n=1):
    return None if v is None else round(v, n)


@router.post("/signals")
async def signals(req: StockReq, request: Request, user: User = Depends(current_user)):
    base = await _base(request, user, req)
    items = _signals(base, req.loc)
    counts = {str(k): sum(1 for x in items if x["level"] == k) for k in (0, 1, 2)}
    risk = {str(k): sum(x["risk"] for x in items if x["level"] == k) for k in (0, 1, 2)}
    new = {str(k): sum(1 for x in items if x["level"] == k and x["new"]) for k in (0, 1, 2)}
    if req.level is not None:
        items = [x for x in items if x["level"] == req.level]
    items.sort(key=lambda x: (-x["risk"], x["days"] if x["days"] is not None else 1e9))
    page = items[req.offset:req.offset + req.limit]
    for x in page:
        for k in ("days", "prev", "br_days"):
            if k in x:
                x[k] = _r(x[k])
        x["risk"], x["sales_day"] = round(x["risk"]), _r(x["sales_day"], 2)
        x["accel"] = _r(x.get("accel"), 2)
    return {"snap": base["snap"], "counts": counts, "risk": risk, "new": new, "total": len(items), "items": page,
            "has_prev": bool(base["pb"] or base["pw"]), "links": base.get("links")}


# ---------------------------------------------------------------------------
@router.post("/map")
async def stock_map(req: StockReq, request: Request, user: User = Depends(current_user)):
    base = await _base(request, user, req)
    meta = base["meta"]
    br = meta["branches"]
    cat_ids = list(meta["cats"])
    cat = req.map_category if req.map_category in meta["cats"] else None
    if cat is None:   # eng ko'p sotilgan kategoriya
        tot: dict[int, float] = {}
        for r in base["b"]:
            tot[r[1]] = tot.get(r[1], 0) + r[6]
        cat = max(tot, key=tot.get) if tot else (cat_ids[0] if cat_ids else None)
    if cat is None:
        return {"snap": base["snap"], "category": None, "categories": [], "brands": [], "rows": []}
    # brendlar — savdo bo'yicha TOP 8
    bs: dict[str, float] = {}
    for r in base["b"]:
        if r[1] == cat:
            bs[r[2]] = bs.get(r[2], 0) + r[6]
    brands = [b for b, _ in sorted(bs.items(), key=lambda x: -x[1])[:8]]
    cells: dict = {}       # (row_id) → {brand: [sq, q]}
    names: dict = {}
    if req.rows == "w":
        for wid, c, brand, sq, sv, q, a, oos, q7 in base["w"]:
            if c == cat:
                x = cells.setdefault(wid, {}).setdefault(brand, [0.0, 0.0, 0.0])
                x[0] += sq
                x[1] += q
                x[2] += q7
                names[wid] = meta["wh"].get(wid)
    else:
        for b, c, brand, sq, sv, q, a, oos, q7 in base["b"]:
            if c != cat:
                continue
            bi = br.get(b, {})
            if req.rows == "b" and req.wid and bi.get("warehouse_id") != req.wid:
                continue
            key = bi.get("region_id") if req.rows == "region" else b
            names[key] = meta["regions"].get(key) if req.rows == "region" else bi.get("name")
            x = cells.setdefault(key, {}).setdefault(brand, [0.0, 0.0, 0.0])
            x[0] += sq
            x[1] += q
            x[2] += q7
    rows, tot = [], {}
    for key, cs in cells.items():
        row = {"id": key, "name": names.get(key), "cells": {}, "qty": {b: round(cs[b][0]) for b in brands if b in cs}}
        s_all = q_all = 0.0
        for b in brands:
            v = cs.get(b)
            row["cells"][b] = None if not v else (_r(_days_eff(*v)) if v[1] > 0 else (-1 if v[0] > 0 else None))
            if v:
                t = tot.setdefault(b, [0.0, 0.0, 0.0])
                t[0] += v[0]
                t[1] += v[1]
                t[2] += v[2]
        q7_all = 0.0
        for v in cs.values():
            s_all += v[0]
            q_all += v[1]
            q7_all += v[2]
        row["all"] = _r(_days_eff(s_all, q_all, q7_all))
        row["worst"] = min([d for d in row["cells"].values() if d is not None and d >= 0] or [1e9])
        rows.append(row)
    rows.sort(key=lambda r: (r["worst"], r["name"] or ""))
    jami = {"cells": {b: (_r(_days_eff(*tot[b])) if b in tot and tot[b][1] > 0 else None) for b in brands},
            "qty": {b: round(tot[b][0]) for b in brands if b in tot}}
    return {"snap": base["snap"], "category": cat, "category_name": meta["cats"].get(cat),
            "categories": [{"id": k, "name": v} for k, v in sorted(meta["cats"].items(), key=lambda x: x[1])],
            "brands": brands, "rows": rows, "total": jami, "rows_kind": req.rows, "wid": req.wid,
            "wh_name": meta["wh"].get(req.wid) if req.wid else None}


# ---------------------------------------------------------------------------
@router.post("/overview")
async def overview(req: StockReq, request: Request, user: User = Depends(current_user)):
    base = await _base(request, user, req)
    meta, br, rate = base["meta"], base["meta"]["branches"], base.get("rate") or 0
    reg: dict = {}
    wh_b: dict = {}
    tb = {"sq": 0.0, "sv": 0.0, "q": 0.0, "a": 0.0, "oos": 0}
    for b, c, brand, sq, sv, q, a, oos, q7 in base["b"]:
        bi = br.get(b, {})
        for d_, k in ((reg, bi.get("region_id")), (wh_b, bi.get("warehouse_id"))):
            x = d_.setdefault(k, {"sq": 0.0, "sv": 0.0, "q": 0.0, "a": 0.0, "oos": 0, "br": set()})
            x["sq"] += sq
            x["sv"] += sv
            x["q"] += q
            x["a"] += a
            x["oos"] += oos
            x["br"].add(b)
        tb["sq"] += sq
        tb["sv"] += sv
        tb["q"] += q
        tb["a"] += a
        tb["oos"] += oos
    tw = {"sq": 0.0, "sv": 0.0, "q": 0.0}
    whs: dict = {}
    for wid, c, brand, sq, sv, q, a, oos, q7 in base["w"]:
        x = whs.setdefault(wid, {"sq": 0.0, "sv": 0.0, "q": 0.0})
        x["sq"] += sq
        x["sv"] += sv
        x["q"] += q
        tw["sq"] += sq
        tw["sv"] += sv
        tw["q"] += q
    sig = _signals(base, "all")
    counts = {str(k): sum(1 for x in sig if x["level"] == k) for k in (0, 1, 2)}
    regions = [{"id": k, "name": meta["regions"].get(k), "stock": v["sq"], "value": v["sv"], "days": _r(_days(v["sq"], v["q"])),
                "oos": v["oos"], "branches": len(v["br"]), "sales_day": v["a"] / WINDOW}
               for k, v in reg.items() if k is not None]
    regions.sort(key=lambda r: (r["days"] if r["days"] is not None else 1e9))
    warehouses = []
    for wid in sorted(set(whs) | {k for k in wh_b if k}, key=lambda k: meta["wh"].get(k) or ""):
        w_ = whs.get(wid, {"sq": 0.0, "sv": 0.0, "q": 0.0})
        b_ = wh_b.get(wid, {"sq": 0.0, "q": 0.0, "oos": 0, "br": set()})
        q = w_["q"] or b_["q"]
        warehouses.append({"id": wid, "name": meta["wh"].get(wid), "stock": w_["sq"], "value": w_["sv"] * rate,
                           "days": _r(_days(w_["sq"], q)), "br_stock": b_["sq"], "br_days": _r(_days(b_["sq"], b_["q"])),
                           "branches": len(b_["br"]), "oos": b_["oos"],
                           "regions": sorted({meta["regions"].get(br[b]["region_id"]) for b in b_["br"] if b in br})})
    q_all = tb["q"]
    lsum = (await _lists(request, user, req)).get("summary")
    return {"snap": base["snap"], "rate": rate, "counts": counts, "lists": lsum, "links": base.get("links"),
            "total": {"br_stock": tb["sq"], "br_value": tb["sv"], "wh_stock": tw["sq"], "wh_value": tw["sv"] * rate,
                      "days": _r(_days(tb["sq"] + tw["sq"], q_all)), "br_days": _r(_days(tb["sq"], q_all)),
                      "wh_days": _r(_days(tw["sq"], q_all)), "oos": tb["oos"], "sales_day": tb["a"] / WINDOW},
            "warehouses": warehouses, "regions": regions}


# ---------------------------------------------------------------------------
@router.post("/detail")
async def detail(req: StockReq, request: Request, user: User = Depends(current_user)):
    """Joy × kategoriya × brend: holat, qoldiq tarixi (kunga yetadi), SKU'lar, boshqa omborlardagi qoldiq."""
    async def run():
        pool = request.app.state.pool
        async with pool.acquire() as con:
            cats, attrs = await _ctx(con, req, user)
            sn = await snapshots(con, req.date_to)
            ls = sn["sales_to"]
            f = req.model_copy()
            if req.map_category:
                if req.map_category not in cats:
                    raise HTTPException(403, "Kategoriya ochiq emas")
                cats = [req.map_category]
            if req.brand:
                f = f.model_copy(update={"brands": [req.brand]})
            # joy filtri
            if req.loc_type == "b" and req.loc_id:
                f = f.model_copy(update={"branch_ids": [req.loc_id]})
            elif req.loc_type == "region" and req.loc_id:
                f = f.model_copy(update={"region_ids": [req.loc_id]})
            wid = None
            if req.loc_type == "w" and req.loc_id:
                wid = req.loc_id
            elif req.loc_type == "b" and req.loc_id:
                wid = await con.fetchval("SELECT warehouse_id FROM branches WHERE id = $1", req.loc_id)
            s = _SQL(f, cats, attrs)
            w = s.ctes()
            if wid and req.loc_type == "w":
                w += f", bw AS (SELECT * FROM bb WHERE warehouse_id = {s.p(wid)})"
            else:
                w += ", bw AS (SELECT * FROM bb)"
            D = s.p(sn["b"])
            E = s.p(_end(sn["b"], ls) or req.date_to)
            DW = s.p(sn["w"])
            WID = s.p(wid)
            skus = await con.fetch(f"""
WITH {w},
cov AS (SELECT DISTINCT branch_id FROM stock_branch WHERE snap_date = {D}),
bst AS (SELECT s.product_id, s.branch_id, s.qty FROM stock_branch s JOIN bw USING (branch_id) JOIN pd USING (product_id)
        WHERE s.snap_date = {D}),
sl AS (SELECT e.product_id, e.branch_id, sum(e.qty) q, sum(e.amount) a,
              sum(e.qty) FILTER (WHERE e.sale_date > {E}::date - 7) q7 FROM sales_fact e JOIN bw USING (branch_id)
       JOIN pd USING (product_id) JOIN cov USING (branch_id)
       WHERE e.sale_date BETWEEN {E}::date - {WINDOW - 1} AND {E} GROUP BY 1, 2),
j AS (SELECT product_id, branch_id, COALESCE(bst.qty, 0) sq, COALESCE(sl.q, 0) q, COALESCE(sl.a, 0) a,
             COALESCE(sl.q7, 0) q7
      FROM bst FULL JOIN sl USING (product_id, branch_id)),
w AS (SELECT product_id, sum(qty) wq FROM stock_wh WHERE snap_date = {DW} AND warehouse_id = {WID}::int GROUP BY 1),
oth AS (SELECT s.product_id, s.warehouse_id, s.qty FROM stock_wh s
        WHERE s.snap_date = {DW} AND s.warehouse_id IS DISTINCT FROM {WID}::int AND s.qty > 0)
SELECT pd.product_id, pd.pname, pd.brand, coalesce(sum(j.sq), 0)::float8 bq, coalesce(sum(j.q), 0)::float8 q,
       coalesce(sum(j.a), 0)::float8 a, count(*) FILTER (WHERE j.sq <= 0 AND j.q > 0) oos,
       coalesce(sum(j.q7), 0)::float8 q7,
       coalesce(max(w.wq), 0)::float8 wq,
       (SELECT json_agg(json_build_object('wid', o.warehouse_id, 'qty', o.qty) ORDER BY o.qty DESC) FROM oth o
        WHERE o.product_id = pd.product_id) other
FROM pd LEFT JOIN j USING (product_id) LEFT JOIN w USING (product_id)
GROUP BY pd.product_id, pd.pname, pd.brand
HAVING coalesce(sum(j.sq), 0) > 0 OR coalesce(sum(j.q), 0) > 0 OR coalesce(max(w.wq), 0) > 0""", *s.params)
            whn = {r["id"]: r["name"] for r in await con.fetch("SELECT id, name FROM warehouses")}
            items = []
            tot = {"bq": 0.0, "q": 0.0, "a": 0.0, "wq": 0.0, "oos": 0, "q7": 0.0}
            for r in skus:
                tot["bq"] += r["bq"]
                tot["q"] += r["q"]
                tot["a"] += r["a"]
                tot["wq"] += r["wq"]
                tot["oos"] += r["oos"]
                tot["q7"] += r["q7"]
                own = r["wq"] if req.loc_type == "w" else r["bq"]
                d = _days_eff(own, r["q"], r["q7"])
                other = json.loads(r["other"]) if r["other"] else []
                items.append({"id": r["product_id"], "name": r["pname"], "brand": r["brand"], "stock": own,
                              "br_stock": r["bq"], "wh_stock": r["wq"], "days": _r(d), "oos": r["oos"],
                              "risk": round(r["a"] / WINDOW), "sales_day": _r(r["q"] / WINDOW, 2),
                              "level": _level(d, _accel(r["q"], r["q7"])),
                              "other": [{"name": whn.get(o["wid"]), "qty": o["qty"]} for o in other[:3]]})
            items.sort(key=lambda x: (x["days"] if x["days"] is not None else 1e9, -x["risk"]))
            # kunlik savdo (30 kun) — sur'at grafigi uchun (bitta qoldiq sanasi bilan ham ishlaydi)
            s3 = _SQL(f, cats, attrs)
            w3 = s3.ctes()
            if wid and req.loc_type == "w":
                w3 += f", bw AS (SELECT * FROM bb WHERE warehouse_id = {s3.p(wid)})"
            else:
                w3 += ", bw AS (SELECT * FROM bb)"
            E3 = s3.p(_end(sn["b"] or sn["w"], ls) or req.date_to)
            daily = [{"date": r["d"].isoformat(), "qty": r["q"]} for r in await con.fetch(f"""
WITH {w3}
SELECT e.sale_date d, sum(e.qty)::float8 q FROM sales_fact e JOIN bw USING (branch_id) JOIN pd USING (product_id)
WHERE e.sale_date BETWEEN {E3}::date - {WINDOW - 1} AND {E3} GROUP BY 1 ORDER BY 1""", *s3.params)]
            links = await con.fetchrow("SELECT count(*) total, count(*) FILTER (WHERE warehouse_id IS NULL) unlinked FROM branches")
            # tarix: har bir qoldiq sanasi bo'yicha kunga yetadi
            tbl = "stock_wh" if req.loc_type == "w" else "stock_branch"
            dates = [r[0] for r in await con.fetch(
                f"SELECT DISTINCT snap_date FROM {tbl} WHERE snap_date <= $1 ORDER BY 1 DESC LIMIT 14",
                sn["w"] if req.loc_type == "w" else sn["b"])] if (sn["w"] if req.loc_type == "w" else sn["b"]) else []
            hist = []
            for d0 in sorted(dates):
                s2 = _SQL(f, cats, attrs)
                w2 = s2.ctes()
                if wid and req.loc_type == "w":
                    w2 += f", bw AS (SELECT * FROM bb WHERE warehouse_id = {s2.p(wid)})"
                else:
                    w2 += ", bw AS (SELECT * FROM bb)"
                D2, E2 = s2.p(d0), s2.p(_end(d0, ls))
                if req.loc_type == "w":
                    stock_sql = (f"(SELECT coalesce(sum(s.qty), 0) FROM stock_wh s JOIN pd USING (product_id) "
                                 f"WHERE s.snap_date = {D2} AND s.warehouse_id = {s2.p(wid)}::int)")
                else:
                    stock_sql = (f"(SELECT coalesce(sum(s.qty), 0) FROM stock_branch s JOIN bw USING (branch_id) "
                                 f"JOIN pd USING (product_id) WHERE s.snap_date = {D2})")
                r = await con.fetchrow(f"""
WITH {w2}
SELECT {stock_sql}::float8 sq,
       (SELECT coalesce(sum(e.qty), 0) FROM sales_fact e JOIN bw USING (branch_id) JOIN pd USING (product_id)
        WHERE e.sale_date BETWEEN {E2}::date - {WINDOW - 1} AND {E2})::float8 q,
       (SELECT coalesce(sum(e.qty), 0) FROM sales_fact e JOIN bw USING (branch_id) JOIN pd USING (product_id)
        WHERE e.sale_date BETWEEN {E2}::date - 6 AND {E2})::float8 q7""", *s2.params)
                hist.append({"date": d0.isoformat(), "stock": r["sq"], "days": _r(_days_eff(r["sq"], r["q"], r["q7"]))})
            name = None
            if req.loc_type == "w":
                name = whn.get(wid)
            elif req.loc_type == "b":
                name = await con.fetchval("SELECT name FROM branches WHERE id = $1", req.loc_id)
            elif req.loc_type == "region":
                name = await con.fetchval("SELECT name FROM regions WHERE id = $1", req.loc_id)
            nb = await con.fetchval("SELECT count(*) FROM branches WHERE warehouse_id = $1", wid) if wid and req.loc_type == "w" else None
            own = tot["wq"] if req.loc_type == "w" else tot["bq"]
            d = _days_eff(own, tot["q"], tot["q7"])
            ac = _accel(tot["q"], tot["q7"])
            return {"snap": {k: (v.isoformat() if v else None) for k, v in sn.items()},
                    "loc_type": req.loc_type, "loc_id": req.loc_id, "name": name or "Jami", "brand": req.brand,
                    "category": await con.fetchval("SELECT name FROM categories WHERE id = $1", cats[0]) if len(cats) == 1 else None,
                    "branches": nb, "level": _level(d, ac), "accel": _r(ac, 2),
                    "total": {"stock": own, "br_stock": tot["bq"], "wh_stock": tot["wq"], "days": _r(d),
                              "br_days": _r(_days(tot["bq"], tot["q"])), "sales_day": _r(tot["q"] / WINDOW, 2),
                              "sales_day7": _r(tot["q7"] / 7, 2),
                              "risk": round(tot["a"] / WINDOW), "oos": tot["oos"]},
                    "history": hist, "daily": daily, "sales_to": ls.isoformat() if ls else None,
                    "links": dict(links), "items": items[:60]}
    payload = req.model_dump(mode="json", include=set(Filters.model_fields) | {"loc_type", "loc_id", "brand", "map_category"})
    return await _M().cached(request, user, "stock_detail", payload, run)


# ---------------------------------------------------------------------------
BY_DIMS = {"category": "pd.category_id::text", "brand": "pd.brand", "sku": "pd.product_id::text",
           "region": "bb.region_id::text", "cluster": "bb.cluster_id::text", "branch": "bb.branch_id::text",
           "all": "'all'"}


@router.post("/by")
async def stock_by(req: StockReq, request: Request, user: User = Depends(current_user)):
    """O'lchov bo'yicha qoldiq: {key: {stock, wh, days, oos, value}}. Joy o'lchovlarida — faqat filial qoldig'i;
    tovar o'lchovlarida va «all» da — filial + ombor qoldig'i (kunga yetadi — tarmoq bo'yicha)."""
    if req.dim not in BY_DIMS:
        return {"snap": None, "items": {}}

    async def run():
        pool = request.app.state.pool
        async with pool.acquire() as con:
            cats, attrs = await _ctx(con, req, user)
            sn = await snapshots(con, req.date_to)
            if not sn["b"] and not sn["w"]:
                return {"snap": None, "items": {}}
            ls = sn["sales_to"]
            key = BY_DIMS[req.dim]
            loc_dim = req.dim in ("region", "cluster", "branch")
            s = _SQL(req, cats, attrs)
            w = s.ctes()
            D = s.p(sn["b"] or sn["w"])
            E = s.p(_end(sn["b"] or sn["w"], ls))
            DW = None if loc_dim else s.p(sn["w"] or sn["b"])
            wh_part = "" if loc_dim else f"""
, wst AS (SELECT {key} k, sum(s.qty) wq, sum(s.value_usd) wv FROM stock_wh s JOIN pd USING (product_id)
          WHERE s.snap_date = {DW} AND s.warehouse_id IN (SELECT DISTINCT warehouse_id FROM bb) GROUP BY 1)"""
            rows = await con.fetch(f"""
WITH {w},
cov AS (SELECT DISTINCT branch_id FROM stock_branch WHERE snap_date = {D}),
st AS (SELECT s.branch_id, s.product_id, s.qty, s.cost_uzs, s.price_uzs FROM stock_branch s JOIN bb USING (branch_id)
       JOIN pd USING (product_id) WHERE s.snap_date = {D}),
sl AS (SELECT e.branch_id, e.product_id, sum(e.qty) q FROM sales_fact e JOIN bb USING (branch_id)
       JOIN pd USING (product_id) WHERE e.sale_date BETWEEN {E}::date - {WINDOW - 1} AND {E} GROUP BY 1, 2),
j AS (SELECT branch_id, product_id, COALESCE(st.qty, 0) sq, COALESCE(st.cost_uzs, 0) sv,
             COALESCE(st.price_uzs, 0) sp, COALESCE(sl.q, 0) q, (branch_id IN (SELECT branch_id FROM cov)) cov
      FROM st FULL JOIN sl USING (branch_id, product_id)),
agg AS (SELECT {key} k, sum(sq) sq, sum(sv) sv, sum(sp) sp, sum(q) q, count(*) FILTER (WHERE cov AND sq <= 0 AND q > 0) oos
        FROM j JOIN pd USING (product_id) JOIN bb USING (branch_id) GROUP BY 1){wh_part}
SELECT agg.k, agg.sq::float8, agg.sv::float8, agg.sp::float8, agg.q::float8, agg.oos{', coalesce(wst.wq, 0)::float8 wq, coalesce(wst.wv, 0)::float8 wv' if not loc_dim else ', 0::float8 wq, 0::float8 wv'}
FROM agg {'LEFT JOIN wst ON wst.k = agg.k' if not loc_dim else ''}""", *s.params)
            rate = float(await con.fetchval("SELECT rate FROM fx_periods WHERE $1 >= valid_from AND $1 < valid_to",
                                            sn["w"] or sn["b"]) or 0)
            items = {}
            for r in rows:
                stock = r["sq"] + r["wq"]
                d = _days(stock, r["q"])
                if stock <= 0 and r["q"] <= 0:
                    continue
                items[r["k"]] = {"stock": stock, "br": r["sq"], "wh": r["wq"], "days": _r(d), "oos": r["oos"],
                                 "value": r["sv"] + r["wv"] * rate, "sell": r["sp"] + r["wv"] * rate,
                                 "lv": 0 if d is not None and d < LV_DAYS[0] else 1 if d is not None and d < LV_DAYS[1]
                                 else 3 if d is not None and d > 90 else 2 if d is not None and d < LV_DAYS[2] else 4}
            return {"snap": {k: (v.isoformat() if v else None) for k, v in sn.items()}, "items": items}
    payload = req.model_dump(mode="json", include=set(Filters.model_fields) | {"dim"})
    return await _M().cached(request, user, "stock_by", payload, run)


# ---------------------------------------------------------------------------
# Ro'yxatlar: 🚚 ombordan yuborish · 📈 ortiqcha · 💤 harakatsiz
#   Me'yor (yuborish uchun) = 2 oylik savdo = 2 × 30 kunlik savdo (Sanjarbek tanlovi, «savdo × 2»).
#   Yuborish: filialda < 7 kunga yetadigan SKU (30 va 7 kunlik sur'atdan kami), uning omborida bor → me'yor − bor (ombordagidan oshmaydi;
#   bir omborning qoldig'i eng kam qolgan filiallarga birinchi taqsimlanadi).
#   Ortiqcha / harakatsiz — «zona» (ombor + uning filiallari) × SKU: > 90 kunga yetadi / 60 kunda sotuv yo'q.
NORM_MONTHS = 2
SEND_DAYS = 7
OVER_DAYS = 90
IDLE_DAYS = 60


async def _sku_level(con, f: Filters, cats: list[int], attrs: dict, sn: dict):
    """Filial × SKU (qoldiq, tannarx, 30/60 kunlik savdo) va ombor × SKU qoldig'i; SKU ma'lumotlari."""
    ls = sn["sales_to"]
    br_rows, wh_rows = [], []
    s = _SQL(f, cats, attrs)
    w = s.ctes()
    if sn["b"]:
        D, E = s.p(sn["b"]), s.p(_end(sn["b"], ls))
        br_rows = await con.fetch(f"""
WITH {w},
cov AS (SELECT DISTINCT branch_id FROM stock_branch WHERE snap_date = {D}),
st AS (SELECT s.branch_id, s.product_id, s.qty, s.cost_uzs FROM stock_branch s JOIN bb USING (branch_id)
       JOIN pd USING (product_id) WHERE s.snap_date = {D}),
sl AS (SELECT e.branch_id, e.product_id,
              sum(e.qty) FILTER (WHERE e.sale_date > {E}::date - {WINDOW}) q30,
              sum(e.amount) FILTER (WHERE e.sale_date > {E}::date - {WINDOW}) a30, sum(e.qty) q60,
              sum(e.qty) FILTER (WHERE e.sale_date > {E}::date - 7) q7
       FROM sales_fact e JOIN bb USING (branch_id) JOIN pd USING (product_id) JOIN cov USING (branch_id)
       WHERE e.sale_date BETWEEN {E}::date - {IDLE_DAYS - 1} AND {E} GROUP BY 1, 2)
SELECT branch_id, product_id, COALESCE(st.qty, 0)::float8 sq, COALESCE(st.cost_uzs, 0)::float8 sv,
       COALESCE(sl.q30, 0)::float8 q30, COALESCE(sl.a30, 0)::float8 a30, COALESCE(sl.q60, 0)::float8 q60,
       COALESCE(sl.q7, 0)::float8 q7
FROM st FULL JOIN sl USING (branch_id, product_id)""", *s.params)
    if sn["w"]:
        s2 = _SQL(f, cats, attrs)
        w2 = s2.ctes()
        DW = s2.p(sn["w"])
        wh_rows = await con.fetch(f"""
WITH {w2}
SELECT s.warehouse_id, s.product_id, s.qty::float8 wq, COALESCE(s.value_usd, 0)::float8 wv
FROM stock_wh s JOIN pd USING (product_id)
WHERE s.snap_date = {DW} AND s.qty > 0 AND s.warehouse_id IN (SELECT warehouse_id FROM bb)""", *s2.params)
    pids = list({r["product_id"] for r in br_rows} | {r["product_id"] for r in wh_rows})
    prod = {r["product_id"]: dict(r) for r in await con.fetch(
        """SELECT p.product_id, p.name, p.brand, c.name cat FROM products p JOIN categories c ON c.id = p.category_id
           WHERE p.product_id = ANY($1::bigint[])""", pids)}
    return br_rows, wh_rows, prod


def build_lists(br_rows, wh_rows, prod: dict, meta: dict, rate: float) -> dict:
    br, whn = meta["branches"], meta["wh"]
    avail: dict[tuple, float] = {(r["warehouse_id"], r["product_id"]): r["wq"] for r in wh_rows}
    # --- ombordan yuborish
    cand = []
    for r in br_rows:
        q30, sq = r["q30"], r["sq"]
        if q30 <= 0:
            continue
        days = _days_eff(sq, q30, r["q7"])
        if days >= SEND_DAYS:
            continue
        bi = br.get(r["branch_id"], {})
        wid = bi.get("warehouse_id")
        if not wid or avail.get((wid, r["product_id"]), 0) <= 0:
            continue
        norm = math.ceil(NORM_MONTHS * q30)
        need = norm - sq
        if need >= 1:
            cand.append((days, -need, r, wid, norm, need))
    cand.sort(key=lambda x: (x[0], x[1]))
    groups: dict[int, dict] = {}
    left = dict(avail)
    for days, _, r, wid, norm, need in cand:
        have = left.get((wid, r["product_id"]), 0)
        send = int(min(need, have))
        if send < 1:
            continue
        left[(wid, r["product_id"])] = have - send
        p = prod.get(r["product_id"], {})
        bi = br.get(r["branch_id"], {})
        g = groups.setdefault(wid, {"id": wid, "name": whn.get(wid), "positions": 0, "qty": 0, "items": []})
        g["positions"] += 1
        g["qty"] += send
        g["items"].append({"branch": bi.get("name"), "region": bi.get("region"), "pid": r["product_id"],
                           "name": p.get("name"), "brand": p.get("brand"), "cat": p.get("cat"),
                           "have": r["sq"], "norm": norm, "send": send, "wh": avail[(wid, r["product_id"])],
                           "days": round(days, 1), "risk": round(r["a30"] / WINDOW), "sd": round(q30 / WINDOW, 2)})
    send_groups = sorted(groups.values(), key=lambda g: -g["qty"])
    # --- zona (ombor + filiallari) × SKU
    zone: dict[tuple, dict] = {}
    for r in br_rows:
        wid = br.get(r["branch_id"], {}).get("warehouse_id")
        z = zone.setdefault((wid, r["product_id"]), {"bq": 0.0, "bv": 0.0, "q30": 0.0, "q60": 0.0, "wq": 0.0, "wv": 0.0, "top": []})
        z["bq"] += r["sq"]
        z["bv"] += r["sv"]
        z["q30"] += r["q30"]
        z["q60"] += r["q60"]
        if r["sq"] > 0:
            z["top"].append((r["sq"], br.get(r["branch_id"], {}).get("name")))
    for r in wh_rows:
        z = zone.setdefault((r["warehouse_id"], r["product_id"]), {"bq": 0.0, "bv": 0.0, "q30": 0.0, "q60": 0.0, "wq": 0.0, "wv": 0.0, "top": []})
        z["wq"] += r["wq"]
        z["wv"] += r["wv"] * rate
    over, idle = [], []
    for (wid, pid), z in zone.items():
        stock = z["bq"] + z["wq"]
        if stock <= 0:
            continue
        p = prod.get(pid, {})
        top = sorted(z["top"], key=lambda x: -x[0])
        item = {"pid": pid, "name": p.get("name"), "brand": p.get("brand"), "cat": p.get("cat"),
                "zone": whn.get(wid) if wid else "ombori ko'rsatilmagan", "wid": wid,
                "stock": stock, "br_stock": z["bq"], "wh_stock": z["wq"], "branches": len(top),
                "value": z["bv"] + z["wv"], "top": [{"name": n, "qty": q} for q, n in top[:4]]}
        if z["q60"] <= 0:
            idle.append(item)
        elif z["q30"] > 0:
            d = stock / (z["q30"] / WINDOW)
            if d > OVER_DAYS:
                over.append({**item, "days": round(d)})
    over.sort(key=lambda x: -x["value"])
    idle.sort(key=lambda x: -x["value"])
    return {
        "send": send_groups, "over": over, "idle": idle,
        "summary": {"send": {"positions": sum(g["positions"] for g in send_groups), "qty": sum(g["qty"] for g in send_groups),
                             "warehouses": len(send_groups)},
                    "over": {"positions": len(over), "value": sum(x["value"] for x in over), "qty": sum(x["stock"] for x in over)},
                    "idle": {"positions": len(idle), "value": sum(x["value"] for x in idle), "qty": sum(x["stock"] for x in idle)}},
        "norm": f"{NORM_MONTHS} oylik savdo", "send_days": SEND_DAYS, "over_days": OVER_DAYS, "idle_days": IDLE_DAYS,
    }


async def lists_data(con, f: Filters, cats: list[int], attrs: dict) -> dict:
    sn = await snapshots(con, f.date_to)
    if not sn["b"] and not sn["w"]:
        return {"snap": None}
    meta = {
        "branches": {r["id"]: dict(r) for r in await con.fetch(
            """SELECT b.id, b.name, b.warehouse_id, rg.name region FROM branches b JOIN regions rg ON rg.id = b.region_id""")},
        "wh": {r["id"]: r["name"] for r in await con.fetch("SELECT id, name FROM warehouses")},
    }
    rate = float(await con.fetchval("SELECT rate FROM fx_periods WHERE $1 >= valid_from AND $1 < valid_to",
                                    sn["w"] or sn["b"]) or 0)
    br_rows, wh_rows, prod = await _sku_level(con, f, cats, attrs, sn)
    out = build_lists(br_rows, wh_rows, prod, meta, rate)
    out["snap"] = {k: (v.isoformat() if v else None) for k, v in sn.items()}
    out["rate"] = rate
    return out


async def _lists(request: Request, user: User, f: Filters) -> dict:
    async def run():
        async with request.app.state.pool.acquire() as con:
            cats, attrs = await _ctx(con, f, user)
            return await lists_data(con, f, cats, attrs)
    return await _M().cached(request, user, "stock_lists", f.model_dump(mode="json", include=set(Filters.model_fields)), run)


class ListReq(StockReq):
    kind: Literal["send", "over", "idle"] = "send"


@router.post("/lists")
async def stock_lists(req: ListReq, request: Request, user: User = Depends(current_user)):
    d = await _lists(request, user, req)
    if not d.get("snap"):
        return {"snap": None, "summary": None, "items": [], "groups": [], "total": 0}
    base = {k: d[k] for k in ("snap", "summary", "norm", "send_days", "over_days", "idle_days", "rate")}
    if req.kind == "send":
        return {**base, "groups": [{**g, "items": g["items"][:req.limit]} for g in d["send"]], "total": len(d["send"])}
    items = d[req.kind]
    return {**base, "items": items[req.offset:req.offset + req.limit], "total": len(items)}


# ---------------------------------------------------------------------------
# Excel: signallar, SKU'lar (filial × SKU, < 7 kun), ombordan yuborish, ortiqcha, harakatsiz → botga
LV_NAMES = ["Tugadi", "Tugayapti", "Kamaymoqda"]


def build_xlsx(base: dict, lists: dict, sku_rows: list, prod: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    head_fill, head_font = PatternFill("solid", fgColor="1F6E43"), Font(bold=True, color="FFFFFF")
    fills = [PatternFill("solid", fgColor=c) for c in ("FCE8E8", "FFE7D6", "FFF3D6")]

    def sheet(ws, headers, rows, widths, lv_col=None):
        ws.append(headers)
        for c in ws[1]:
            c.fill, c.font, c.alignment = head_fill, head_font, Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append(r)
            if lv_col is not None and isinstance(r[lv_col], str) and r[lv_col] in LV_NAMES:
                for c in ws[ws.max_row]:
                    c.fill = fills[LV_NAMES.index(r[lv_col])]
        for i, w_ in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w_
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    sig = sorted(_signals(base, "all"), key=lambda x: (x["level"], -x["risk"]))
    ws = wb.active
    ws.title = "Signallar"
    sheet(ws, ["Daraja", "Joy turi", "Joy", "Hudud", "Kategoriya", "Brend", "Qoldiq, dona", "Kunga yetadi",
               "O'tgan hafta, kun", "Filiallarda, dona", "OOS (filial×SKU)", "Savdo, dona/kun", "7 kunlik savdo, ±%",
               "Xavf, so'm/kun", "Yangi"],
          [[LV_NAMES[x["level"]], "Ombor" if x["t"] == "w" else "Filial", x["name"], x.get("sub") or "", x["cat"], x["brand"],
            round(x["stock"]), _r(x["days"]), _r(x["prev"]), round(x["br_stock"]) if x["t"] == "w" else None, x["oos"],
            _r(x["sales_day"], 2), round((x["accel"] - 1) * 100) if x.get("accel") else None,
            round(x["risk"]), "ha" if x["new"] else ""] for x in sig],
          [12, 9, 28, 18, 18, 14, 11, 11, 12, 12, 12, 12, 12, 14, 7], lv_col=0)
    br = base["meta"]["branches"]
    rows = []
    for r in sku_rows:
        if r["q30"] <= 0:
            continue
        d = _days_eff(r["sq"], r["q30"], r["q7"])
        if d >= SEND_DAYS:
            continue
        p, bi = prod.get(r["product_id"], {}), br.get(r["branch_id"], {})
        rows.append([LV_NAMES[0] if d < LV_DAYS[0] else LV_NAMES[1], bi.get("name"), bi.get("region"),
                     base["meta"]["wh"].get(bi.get("warehouse_id"), ""), p.get("cat"), p.get("brand"), r["product_id"],
                     p.get("name"), round(r["sq"]), round(d, 1), round(r["q30"] / WINDOW, 2), round(r["a30"] / WINDOW)])
    rows.sort(key=lambda x: (x[9], -x[11]))
    sheet(wb.create_sheet("SKU'lar"), ["Daraja", "Filial", "Hudud", "Ombor", "Kategoriya", "Brend", "Tovar ID", "Tovar",
                                       "Qoldiq", "Kunga yetadi", "Savdo, dona/kun", "Xavf, so'm/kun"],
          rows, [11, 26, 18, 18, 16, 12, 10, 40, 8, 10, 12, 13], lv_col=0)
    sheet(wb.create_sheet("Ombordan yuborish"), ["Ombor", "Filial", "Hudud", "Kategoriya", "Brend", "Tovar ID", "Tovar",
                                                 "Filialda", f"Me'yor ({lists['norm']})", "Yuborish", "Omborda", "Kunga yetadi"],
          [[g["name"], x["branch"], x["region"], x["cat"], x["brand"], x["pid"], x["name"], round(x["have"]), x["norm"],
            x["send"], round(x["wh"]), x["days"]] for g in lists["send"] for x in g["items"]],
          [18, 26, 18, 16, 12, 10, 40, 9, 14, 10, 10, 11])
    for key, title, extra in (("over", "Ortiqcha", "Kunga yetadi"), ("idle", "Harakatsiz", None)):
        hdr = ["Zona (ombor + filiallari)", "Kategoriya", "Brend", "Tovar ID", "Tovar", "Jami qoldiq", "Filiallarda",
               "Omborda", "Filiallar soni", "Qiymati, so'm", "Eng ko'p"] + ([extra] if extra else [])
        sheet(wb.create_sheet(title), hdr,
              [[x["zone"], x["cat"], x["brand"], x["pid"], x["name"], round(x["stock"]), round(x["br_stock"]), round(x["wh_stock"]),
                x["branches"], round(x["value"]), ", ".join(f"{t['name']} — {round(t['qty'])}" for t in x["top"])]
               + ([x["days"]] if extra else []) for x in lists[key]],
              [22, 16, 12, 10, 40, 10, 10, 10, 10, 15, 40] + ([11] if extra else []))
    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def export_bytes(con, f: Filters, cats: list[int], attrs: dict) -> tuple[bytes, str] | None:
    base = await base_data(con, f, cats, attrs)
    sn = await snapshots(con, f.date_to)
    if not sn["b"] and not sn["w"]:
        return None
    br_rows, wh_rows, prod = await _sku_level(con, f, cats, attrs, sn)
    meta = {"branches": base["meta"]["branches"], "wh": base["meta"]["wh"]}
    lists = build_lists(br_rows, wh_rows, prod, meta, base["rate"])
    d = sn["b"] or sn["w"]
    return build_xlsx(base, lists, br_rows, prod), f"Qoldiq_signallari_{d.strftime('%d.%m.%Y')}.xlsx"


@router.post("/export")
async def stock_export(req: StockReq, request: Request, user: User = Depends(current_user)):
    """Excel faylni botga (foydalanuvchining chatiga) yuboradi."""
    import config
    from aiogram import Bot
    from aiogram.types import BufferedInputFile
    async with request.app.state.pool.acquire() as con:
        cats, attrs = await _ctx(con, req, user)
        res = await export_bytes(con, req, cats, attrs)
    if res is None:
        raise HTTPException(404, "Qoldiq hali yuklanmagan")
    data, fname = res
    if not config.BOT_TOKEN:
        raise HTTPException(503, "Bot sozlanmagan")
    bot = Bot(config.BOT_TOKEN)
    try:
        await bot.send_document(user.id, BufferedInputFile(data, filename=fname),
                                caption="📦 Qoldiq: signallar, SKU'lar, ombordan yuborish, ortiqcha, harakatsiz")
    except Exception as e:  # noqa: BLE001
        raise HTTPException(503, f"Botga yuborib bo'lmadi — botga /start bosilganini tekshiring ({e})")
    finally:
        await bot.session.close()
    return {"ok": True}
